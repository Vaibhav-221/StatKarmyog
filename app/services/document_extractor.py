"""
Document text extractor for the Quiz Generation Engine.

Extracts plain text from uploaded files (.pdf, .pptx, .docx, .txt, .md)
for downstream LLM processing. Enforces minimum content length and
truncates excessively long documents to control token cost.
"""

import logging
from pathlib import Path

from fastapi import UploadFile

logger = logging.getLogger(__name__)

# Supported file extensions (lowercase, with dot)
SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx", ".txt", ".md"}

# Minimum extractable text length (characters)
MIN_TEXT_LENGTH = 200

# Maximum text length sent to the LLM (characters) — prototype cap to
# avoid excessive token cost.  ~12 000 chars ≈ ~3 000 tokens.
MAX_TEXT_LENGTH = 12_000


def _extract_pdf(content: bytes) -> str:
    """Extract text from a PDF file using pypdf."""
    from pypdf import PdfReader
    import io

    reader = PdfReader(io.BytesIO(content))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def _extract_pptx(content: bytes) -> str:
    """Extract text from all slides/shapes in a PowerPoint file."""
    from pptx import Presentation
    import io

    prs = Presentation(io.BytesIO(content))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        texts.append(text)
    return "\n\n".join(texts)


def _extract_docx(content: bytes) -> str:
    """Extract text from all paragraphs in a Word document."""
    from docx import Document
    import io

    doc = Document(io.BytesIO(content))
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    return "\n\n".join(paragraphs)


def _extract_plain(content: bytes) -> str:
    """Decode plain text content (UTF-8)."""
    return content.decode("utf-8", errors="replace")


# Dispatch table: extension -> extractor function
_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".pptx": _extract_pptx,
    ".docx": _extract_docx,
    ".txt": _extract_plain,
    ".md": _extract_plain,
}


async def extract_text(file: UploadFile) -> str:
    """
    Read the uploaded file and extract its text content.

    Raises
    ------
    ValueError
        If the file extension is unsupported or the extracted text is too
        short (< MIN_TEXT_LENGTH characters).

    Returns
    -------
    str
        The extracted (and possibly truncated) text.
    """
    filename = file.filename or ""
    ext = Path(filename).suffix.lower()

    if ext not in _EXTRACTORS:
        raise ValueError(
            f"Unsupported file type '{ext}'. "
            f"Accepted types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    await file.seek(0)
    content = await file.read()
    text = _EXTRACTORS[ext](content)
    text = text.strip()

    if len(text) < MIN_TEXT_LENGTH:
        raise ValueError(
            "Document has no extractable text (or text is too short — "
            f"need at least {MIN_TEXT_LENGTH} characters, got {len(text)})."
        )

    if len(text) > MAX_TEXT_LENGTH:
        logger.warning(
            "Document text is %d chars — truncating to %d chars for LLM input.",
            len(text),
            MAX_TEXT_LENGTH,
        )
        text = text[:MAX_TEXT_LENGTH]

    return text
